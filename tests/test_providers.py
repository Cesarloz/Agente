import asyncio
import json

import httpx
import pytest
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings

from app.providers import ChromaVectorStoreProvider, FAISSVectorStoreProvider, LoaderPipeline
from app.providers.vectors import collection_name
from app.providers.llm import list_models


class LocalEmbeddings(Embeddings):
    """Deterministic tiny embeddings; no model downloads or external requests."""

    def embed_documents(self, texts):
        return [self.embed_query(text) for text in texts]

    def embed_query(self, text):
        normalized = text.lower()
        return [float("horario" in normalized), float("precio" in normalized), 0.1]


def test_collection_separation_survives_slug_collisions_and_model_changes():
    identifiers = [
        collection_name("a/b", "faq", "openai:model-1"),
        collection_name("a-b", "faq", "openai:model-1"),
        collection_name("a/b", "documents", "openai:model-1"),
        collection_name("a/b", "faq", "gemini:model-2"),
    ]
    assert len(set(identifiers)) == 4
    assert all("/" not in item and ".." not in item for item in identifiers)


@pytest.mark.asyncio
async def test_text_loader_normalizes_and_split_keeps_source(tmp_path):
    file = tmp_path / "guide.md"
    file.write_text("# Información\n\n" + ("El horario es de nueve a cinco.\n" * 25), encoding="utf-8")
    pipeline = LoaderPipeline()
    loaded = await pipeline.load(file, ".md")
    chunks = pipeline.split(loaded, chunk_size=150, chunk_overlap=20)
    assert len(chunks) > 1
    assert all(len(chunk.page_content) <= 150 for chunk in chunks)
    assert all(chunk.metadata["source"] == str(file) for chunk in chunks)
    assert [chunk.metadata["chunk_index"] for chunk in chunks] == list(range(len(chunks)))
    assert all("start_index" in chunk.metadata for chunk in chunks)


@pytest.mark.asyncio
async def test_docx_loader_uses_real_file(tmp_path):
    from docx import Document as WordDocument

    path = tmp_path / "guide.docx"
    doc = WordDocument()
    doc.add_paragraph("Horario de atención: 9 a 5.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Lunes"
    table.cell(0, 1).text = "Disponible"
    doc.save(path)
    result = await LoaderPipeline().load(path)
    assert "Horario de atención" in result[0].page_content
    assert "Lunes\tDisponible" in result[0].page_content


@pytest.mark.asyncio
async def test_pdf_loader_extracts_text_and_retains_page_source(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    contents = DecodedStreamObject()
    contents.set_data(b"BT /F1 12 Tf 10 100 Td (Horario de nueve a cinco.) Tj ET")
    page[NameObject("/Contents")] = contents
    path = tmp_path / "guide.pdf"
    with path.open("wb") as stream:
        writer.write(stream)
    result = await LoaderPipeline().load(path)
    assert "Horario de nueve" in result[0].page_content
    assert result[0].metadata == {"source": str(path), "page": 0, "total_pages": 1}


@pytest.mark.asyncio
async def test_loader_rejects_unsupported_format(tmp_path):
    with pytest.raises(ValueError, match="Formato"):
        await LoaderPipeline().load(tmp_path / "program.exe")


@pytest.mark.asyncio
async def test_chroma_persistence_upsert_delete_and_tenant_isolation(tmp_path):
    pytest.importorskip("langchain_chroma")
    embeddings = LocalEmbeddings()
    store = ChromaVectorStoreProvider(persist_dir=tmp_path / "chroma")
    profile = "local-test:three-dimensions"
    await store.upsert("first", "documents", [Document(page_content="horario nueve", metadata={"document_id": "d-1"})], ["chunk-1"], embeddings, profile)
    await store.upsert("second", "documents", [Document(page_content="precio cien")], ["chunk-1"], embeddings, profile)
    result = await store.retrieve("first", "documents", "horario", embeddings, profile)
    assert len(result) == 1
    assert result[0]["metadata"]["agent_id"] == "first"
    assert result[0]["page_content"] == "horario nueve"
    assert await store.retrieve("first", "faq", "horario", embeddings, profile) == []
    assert await store.retrieve("first", "documents", "horario", embeddings, "new-profile") == []
    await store.upsert("first", "documents", [Document(page_content="horario diez")], ["chunk-1"], embeddings, profile)
    restarted = ChromaVectorStoreProvider(persist_dir=tmp_path / "chroma")
    after = await restarted.retrieve("first", "documents", "horario", embeddings, profile)
    assert len(after) == 1
    assert after[0]["page_content"] == "horario diez"
    await restarted.delete("first", "documents", ["chunk-1"], embeddings, profile)
    assert await restarted.retrieve("first", "documents", "horario", embeddings, profile) == []


@pytest.mark.asyncio
async def test_faiss_persistence_uses_json_and_supports_replacement(tmp_path):
    pytest.importorskip("faiss")
    embeddings = LocalEmbeddings()
    store = FAISSVectorStoreProvider(tmp_path)
    await store.upsert("a-1", "faq", [Document(page_content="horario nueve")], ["faq-1"], embeddings, "local")
    await store.upsert("a-1", "faq", [Document(page_content="horario diez")], ["faq-1"], embeddings, "local")
    restarted = FAISSVectorStoreProvider(tmp_path)
    result = await restarted.retrieve("a-1", "faq", "horario", embeddings, "local")
    assert len(result) == 1
    assert result[0]["content"] == "horario diez"
    await restarted.upsert("a-1", "faq", [Document(page_content="precio cien")], ["faq-2"], embeddings, "local")
    relevant = await restarted.retrieve("a-1", "faq", "precio", embeddings, "local", k=1)
    assert relevant[0]["id"] == "faq-2"
    diverse = await restarted.retrieve("a-1", "faq", "precio", embeddings, "local", k=2, search_type="mmr")
    assert {item["id"] for item in diverse} == {"faq-1", "faq-2"}
    assert not list(tmp_path.glob("*.pkl"))
    assert json.loads(next(tmp_path.glob("*.json")).read_text(encoding="utf-8"))["faq-1"]["vector"]
    await restarted.delete("a-1", "faq", ["faq-1", "faq-2"], None, "local")
    assert await restarted.retrieve("a-1", "faq", "horario", embeddings, "local") == []


@pytest.mark.asyncio
async def test_duplicate_upsert_ids_are_rejected_before_embedding(tmp_path):
    store = FAISSVectorStoreProvider(tmp_path)
    with pytest.raises(ValueError, match="únicos"):
        await store.upsert("a", "faq", [Document(page_content="x"), Document(page_content="y")], ["same", "same"], LocalEmbeddings(), "local")


@pytest.mark.asyncio
async def test_faiss_concurrent_provider_instances_preserve_both_writes(tmp_path):
    first, second = FAISSVectorStoreProvider(tmp_path), FAISSVectorStoreProvider(tmp_path)
    await asyncio.gather(
        first.upsert("a", "faq", [Document(page_content="horario")], ["one"], LocalEmbeddings(), "local"),
        second.upsert("a", "faq", [Document(page_content="precio")], ["two"], LocalEmbeddings(), "local"),
    )
    stored = json.loads(next(tmp_path.glob("*.json")).read_text(encoding="utf-8"))
    assert set(stored) == {"one", "two"}


@pytest.mark.asyncio
async def test_gemini_catalog_paginates_and_filters_capabilities(monkeypatch):
    requests = []

    def respond(request):
        requests.append(request)
        assert "key=" not in str(request.url)
        if request.url.params.get("pageToken"):
            return httpx.Response(200, json={"models": [{"name": "models/second-chat", "supportedGenerationMethods": ["generateContent"]}]})
        return httpx.Response(200, json={"models": [
            {"name": "models/first-chat", "supportedGenerationMethods": ["generateContent"]},
            {"name": "models/vector", "supportedGenerationMethods": ["embedContent"]},
        ], "nextPageToken": "page-two"})

    client_class = httpx.AsyncClient
    monkeypatch.setattr("app.providers.llm.httpx.AsyncClient", lambda **kwargs: client_class(transport=httpx.MockTransport(respond), **kwargs))
    result = await list_models("gemini", "unit-test-placeholder")
    assert result == ["first-chat", "second-chat"]
    assert len(requests) == 2
    embeddings = await list_models("gemini", "unit-test-placeholder", purpose="embeddings")
    assert embeddings == ["vector"]


@pytest.mark.asyncio
async def test_model_catalog_propagates_failed_auth_without_fake_models(monkeypatch):
    client_class = httpx.AsyncClient
    transport = httpx.MockTransport(lambda request: httpx.Response(401, json={"error": {"message": "Unauthorized"}}))
    monkeypatch.setattr("app.providers.llm.httpx.AsyncClient", lambda **kwargs: client_class(transport=transport, **kwargs))
    with pytest.raises(httpx.HTTPStatusError):
        await list_models("openai", "unit-test-placeholder")


@pytest.mark.parametrize("provider,model", [("openai", "gpt-5"), ("gemini", "gemini-2.5-flash")])
def test_chat_model_constructor_and_structured_output_need_no_network(monkeypatch, provider, model):
    from pydantic import BaseModel

    from app.providers import GeminiLLMProvider, OpenAILLMProvider

    class ResponseSchema(BaseModel):
        response: str
        answered: bool

    def network_forbidden(*args, **kwargs):
        raise AssertionError("Constructor tests must never make HTTP requests")

    async def async_network_forbidden(*args, **kwargs):
        raise AssertionError("Constructor tests must never make HTTP requests")

    monkeypatch.setattr(httpx.Client, "send", network_forbidden)
    monkeypatch.setattr(httpx.AsyncClient, "send", async_network_forbidden)
    factory = OpenAILLMProvider() if provider == "openai" else GeminiLLMProvider()
    chat = factory.build("obviously-fake-key-for-offline-constructor-test", model)
    structured = chat.with_structured_output(ResponseSchema, include_raw=True)
    assert callable(structured.ainvoke)
